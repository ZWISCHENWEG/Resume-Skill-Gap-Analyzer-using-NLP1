import os
import pytest
import pymupdf
import docx
from src.extraction.document_parser import parse_document
from src.extraction.pdf_parser import extract_text_from_pdf
from src.extraction.docx_parser import extract_text_from_docx

@pytest.fixture
def temp_dir(tmpdir):
    return str(tmpdir)

def create_synthetic_pdf(path, text_pages=None, empty=False):
    doc = pymupdf.open()
    if not empty and text_pages:
        for text in text_pages:
            page = doc.new_page()
            page.insert_text(pymupdf.Point(50, 50), text)
    elif empty:
        doc.new_page() # Just an empty page
    doc.save(path)
    doc.close()
    return path

def create_synthetic_docx(path, paragraphs=None, tables=None, empty=False):
    doc = docx.Document()
    if not empty:
        if paragraphs:
            for p in paragraphs:
                doc.add_paragraph(p)
        if tables:
            for table_data in tables:
                table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                for i, row_data in enumerate(table_data):
                    row = table.rows[i]
                    for j, cell_data in enumerate(row_data):
                        row.cells[j].text = cell_data
    doc.save(path)
    return path

def test_valid_pdf_extraction(temp_dir):
    pdf_path = os.path.join(temp_dir, "valid.pdf")
    create_synthetic_pdf(pdf_path, ["Hello World", "This is a test resume."])
    text = extract_text_from_pdf(pdf_path)
    assert "Hello World" in text
    assert "This is a test resume" in text

def test_empty_pdf(temp_dir):
    pdf_path = os.path.join(temp_dir, "empty.pdf")
    create_synthetic_pdf(pdf_path, empty=True)
    with pytest.raises(ValueError, match="No extractable text found"):
        extract_text_from_pdf(pdf_path)

def test_valid_docx_extraction(temp_dir):
    docx_path = os.path.join(temp_dir, "valid.docx")
    create_synthetic_docx(docx_path, paragraphs=["Hello DOCX", "Resume Content"])
    text = extract_text_from_docx(docx_path)
    assert "Hello DOCX" in text
    assert "Resume Content" in text

def test_docx_with_tables(temp_dir):
    docx_path = os.path.join(temp_dir, "tables.docx")
    create_synthetic_docx(docx_path, tables=[[["Skill", "Years"], ["Python", "5"]]])
    text = extract_text_from_docx(docx_path)
    assert "Skill | Years" in text
    assert "Python | 5" in text

def test_empty_docx(temp_dir):
    docx_path = os.path.join(temp_dir, "empty.docx")
    create_synthetic_docx(docx_path, empty=True)
    with pytest.raises(ValueError, match="No extractable text found"):
        extract_text_from_docx(docx_path)

def test_parse_document_routing(temp_dir):
    pdf_path = os.path.join(temp_dir, "route.pdf")
    docx_path = os.path.join(temp_dir, "route.docx")
    create_synthetic_pdf(pdf_path, ["PDF Content"])
    create_synthetic_docx(docx_path, paragraphs=["DOCX Content"])
    
    pdf_text = parse_document(pdf_path)
    docx_text = parse_document(docx_path)
    
    assert "PDF Content" in pdf_text
    assert "DOCX Content" in docx_text

def test_unsupported_extension(temp_dir):
    txt_path = os.path.join(temp_dir, "test.txt")
    with open(txt_path, "w") as f:
        f.write("Text file")
    
    with pytest.raises(ValueError, match="Unsupported file format"):
        parse_document(txt_path)

def test_file_not_found():
    with pytest.raises(FileNotFoundError):
        parse_document("nonexistent.pdf")

def test_corrupt_pdf(temp_dir):
    corrupt_path = os.path.join(temp_dir, "corrupt.pdf")
    with open(corrupt_path, "w") as f:
        f.write("This is not a real PDF")
    
    with pytest.raises(ValueError, match="Failed to parse PDF"):
        parse_document(corrupt_path)
